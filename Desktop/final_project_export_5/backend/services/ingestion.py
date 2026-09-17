import csv
import io
import re
from typing import List, Dict, Any, Optional
import docx

from backend.models import Question, QuestionType
from backend.config import DOMAINS


VALID_DOMAINS = set(DOMAINS) | {"Research Methodology"}


def parse_docx_questions(file_bytes: bytes, domain: str, phase: int = 1) -> List[Question]:
    if domain not in VALID_DOMAINS:
        raise ValueError(f"Domain '{domain}' is not valid. Must be one of: {', '.join(sorted(VALID_DOMAINS))}")

    document = docx.Document(io.BytesIO(file_bytes))
    answer_key: Dict[int, str] = {}

    # 1. Parse answer key from doc.tables[0] if present (e.g. CS QP)
    if document.tables:
        t = document.tables[0]
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            for i in range(0, len(cells) - 1, 2):
                q_str = cells[i]
                ans_str = cells[i + 1].strip().upper()
                if q_str.isdigit() and ans_str in ("A", "B", "C", "D"):
                    answer_key[int(q_str)] = ans_str

    # 2. Check for paragraph-based answer key (e.g. Research Methodology Set II)
    is_in_para_answer_key = False
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if "ANSWER KEY" in text.upper():
            is_in_para_answer_key = True
            matches = re.findall(r'(\d+)[\.\)]\s*([A-Da-d])', text)
            for q_num, ans in matches:
                answer_key[int(q_num)] = ans.upper()
            continue
        if is_in_para_answer_key:
            matches = re.findall(r'(\d+)[\.\)]\s*([A-Da-d])', text)
            for q_num, ans in matches:
                answer_key[int(q_num)] = ans.upper()

    # 3. Parse questions from paragraphs
    questions_data: List[Question] = []
    current_q_num: Optional[int] = None
    current_question: Optional[str] = None
    current_options: List[str] = []

    q_pattern = re.compile(r'^(\d+)[\.\)]\s+(.*)')
    o_pattern = re.compile(r'^[a-dA-D][\.\)]\s+(.*)')

    def flush_current():
        if current_question:
            q_num = current_q_num if current_q_num is not None else len(questions_data) + 1
            ans_letter = answer_key.get(q_num, "A")
            correct_idx = ord(ans_letter) - ord("A") if ans_letter in ("A", "B", "C", "D") else 0
            opts = current_options[:4] if current_options else ["Option A", "Option B", "Option C", "Option D"]
            while len(opts) < 4:
                opts.append(f"Option {chr(ord('A') + len(opts))}")
            questions_data.append(Question(
                order_no=len(questions_data) + 1,
                domain=domain,
                phase=phase,
                question_type=QuestionType.mcq,
                question_text=current_question,
                options=opts,
                correct_option=correct_idx
            ))

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Stop parsing questions when reaching answer key text section
        if "ANSWER KEY" in text.upper():
            break

        q_match = q_pattern.match(text)
        if q_match:
            flush_current()
            current_q_num = int(q_match.group(1))
            current_question = q_match.group(2).strip()
            current_options = []
        else:
            splits = re.split(r'\s+[A-Da-d][\.\)]\s+', " " + text)
            if len(splits) > 1:
                for s in splits[1:]:
                    opt = s.strip()
                    if opt:
                        current_options.append(opt)
            elif current_question:
                o_match = o_pattern.match(text)
                if o_match:
                    current_options.append(o_match.group(1).strip())
                else:
                    if not current_options:
                        current_question += " " + text
                    else:
                        current_options[-1] += " " + text

    flush_current()
    return questions_data


def parse_csv_questions(file_bytes: bytes, domain: str, phase: int = 1) -> List[Question]:
    if domain not in VALID_DOMAINS:
        raise ValueError(f"Domain '{domain}' is not valid. Must be one of: {', '.join(sorted(VALID_DOMAINS))}")

    text = file_bytes.decode('utf-8-sig')
    reader = csv.DictReader(io.StringIO(text))
    questions_data: List[Question] = []

    for i, row in enumerate(reader):
        row_domain = (row.get("domain") or domain).strip()
        if row_domain not in VALID_DOMAINS:
            raise ValueError(f"Row {i+2}: Domain '{row_domain}' is not valid")

        q_type_str = (row.get("question_type") or "mcq").strip().lower()
        if q_type_str not in ("mcq", "written"):
            raise ValueError(f"Row {i+2}: question_type must be 'mcq' or 'written'")

        question_text = (row.get("question_text") or row.get("Question") or "").strip()
        if not question_text:
            continue

        if q_type_str == "written":
            questions_data.append(Question(
                order_no=len(questions_data) + 1,
                domain=row_domain,
                phase=phase,
                question_type=QuestionType.written,
                question_text=question_text,
                options=None,
                correct_option=None
            ))
            continue

        # MCQ options
        opt_a = (row.get("option_a") or row.get("Option A") or "").strip()
        opt_b = (row.get("option_b") or row.get("Option B") or "").strip()
        opt_c = (row.get("option_c") or row.get("Option C") or "").strip()
        opt_d = (row.get("option_d") or row.get("Option D") or "").strip()
        options = [o for o in [opt_a, opt_b, opt_c, opt_d] if o]
        if len(options) < 2:
            raise ValueError(f"Row {i+2}: MCQ requires at least 2 options")

        correct_ans_str = (row.get("correct_option") or row.get("Correct Answer") or "A").strip().upper()
        correct_idx = 0
        if correct_ans_str in ('A', 'B', 'C', 'D'):
            correct_idx = ord(correct_ans_str) - ord('A')
        else:
            try:
                correct_idx = int(correct_ans_str)
                if correct_idx >= len(options):
                    correct_idx = correct_idx - 1  # 1-based index conversion
            except ValueError:
                correct_idx = 0

        questions_data.append(Question(
            order_no=len(questions_data) + 1,
            domain=row_domain,
            phase=phase,
            question_type=QuestionType.mcq,
            question_text=question_text,
            options=options,
            correct_option=correct_idx
        ))

    return questions_data

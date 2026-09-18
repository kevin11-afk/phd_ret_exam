from datetime import datetime
from typing import List
import csv
import io
import openpyxl

from fastapi import APIRouter, Depends, WebSocket, WebSocketDisconnect, Query, UploadFile, File, Form, HTTPException, status
from sqlalchemy.orm import Session

from backend import models, schemas, exam_engine
from backend.database import get_db, SessionLocal
from backend import security
from backend.config import MAX_STRIKES
from backend.websocket_manager import manager
from backend.services.ingestion import parse_docx_questions, parse_csv_questions, VALID_DOMAINS

router = APIRouter(tags=["admin"])


@router.get("/admin/sessions", response_model=List[schemas.CandidateStatusOut])
def list_sessions(db: Session = Depends(get_db), admin: models.User = Depends(security.get_current_admin)):
    sessions = (
        db.query(models.ExamSession)
        .filter(models.ExamSession.status != models.SessionStatus.not_started)
        .all()
    )
    out = []
    for s in sessions:
        s = exam_engine.enforce_expiry(db, s)
        nav = exam_engine.build_nav(db, s)
        answered = sum(1 for n in nav if n["answered"])
        seconds_ago = int((datetime.utcnow() - s.last_activity).total_seconds())
        out.append(schemas.CandidateStatusOut(
            session_id=s.id,
            user_id=s.user_id,
            full_name=s.user.full_name,
            email=s.user.email,
            status=s.status.value,
            strikes=s.strikes,
            max_strikes=MAX_STRIKES,
            answered_count=answered,
            total_questions=len(s.question_order or []),
            time_remaining_seconds=exam_engine.time_remaining_seconds(s),
            last_activity_seconds_ago=seconds_ago,
            score=s.score,
        ))
    return out


@router.post("/admin/sessions/{session_id}/disqualify")
async def disqualify_candidate_session(
    session_id: int,
    db: Session = Depends(get_db),
    admin: models.User = Depends(security.get_current_admin),
):
    session = db.query(models.ExamSession).filter(models.ExamSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Exam session not found")

    exam_engine.disqualify_session(db, session, reason="Admin manual disqualification")
    nav = exam_engine.build_nav(db, session)
    answered = sum(1 for n in nav if n["answered"])
    await manager.broadcast({
        "type": "candidate_update",
        "session_id": session.id,
        "user_id": session.user_id,
        "full_name": session.user.full_name,
        "email": session.user.email,
        "status": session.status.value,
        "strikes": session.strikes,
        "max_strikes": MAX_STRIKES,
        "answered_count": answered,
        "total_questions": len(session.question_order or []),
        "time_remaining_seconds": exam_engine.time_remaining_seconds(session),
        "score": session.score,
    })
    return {"status": "disqualified", "session_id": session_id}


@router.post("/admin/questions/upload")
async def upload_questions(
    file: UploadFile = File(...),
    domain: str = Form(...),
    phase: int = Form(...),
    db: Session = Depends(get_db),
    admin: models.User = Depends(security.get_current_admin)
):
    if domain not in VALID_DOMAINS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Domain '{domain}' is not in configured DOMAINS: {', '.join(sorted(VALID_DOMAINS))}"
        )
    if phase not in (1, 2):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Phase must be 1 or 2")

    contents = await file.read()
    filename = file.filename.lower()

    try:
        if filename.endswith(".docx"):
            questions = parse_docx_questions(contents, domain, phase)
        elif filename.endswith(".csv"):
            questions = parse_csv_questions(contents, domain, phase)
        else:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported file type (must be .docx or .csv)")

        for q in questions:
            db.add(q)
        db.commit()
        return {"status": "success", "inserted": len(questions)}
    except HTTPException:
        db.rollback()
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.post("/admin/users/upload")
async def upload_users(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin: models.User = Depends(security.get_current_admin)
):
    contents = await file.read()
    try:
        text = contents.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="File must be valid UTF-8 text")

    reader = csv.DictReader(io.StringIO(text))
    inserted = 0
    default_hash = security.hash_password("DefaultPass@2026")

    for row in reader:
        email = row.get("email", "").strip()
        name = row.get("name", "").strip()
        if not email or not name:
            continue

        existing = db.query(models.User).filter(models.User.email == email).first()
        if not existing:
            user = models.User(
                email=email,
                full_name=name,
                hashed_password=default_hash,
                role=models.UserRole.candidate,
                must_reset_password=True,
                is_active=True
            )
            db.add(user)
            inserted += 1

    db.commit()
    return {"status": "success", "inserted": inserted}


@router.post("/admin/users/upload-xlsx", response_model=schemas.BulkImportXlsxResult)
async def upload_users_xlsx(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin: models.User = Depends(security.get_current_admin)
):
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="File must be .xlsx")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid Excel file")

    if "Application Details" not in wb.sheetnames:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Sheet 'Application Details' not found")

    ws = wb["Application Details"]

    def map_program(program):
        p = program.lower()
        if "law" in p: return "Law"
        if "physiology" in p or "microbiology" in p: return "Bio-Science"
        if "computer science" in p: return "Computer Science & Engineering"
        if "management studies" in p or "commerce" in p: return "Commerce & Management Studies"
        if "english" in p: return "English"
        if "mathematics" in p: return "Mathematics"
        if "agronomy" in p or "agricultural sciences" in p: return "Agricultural Sciences"
        if "tamil" in p: return "Tamil"
        return "Unknown"

    created = 0
    already_existed = 0
    errors = []
    default_hash = security.hash_password("DefaultPass@2026")

    # Header is on row 4, data starts at row 5
    for r in range(5, ws.max_row + 1):
        # Read columns: D=First(4), E=Middle(5), F=Last(6), H=Email(8), B=Program(2)
        fname = ws.cell(row=r, column=4).value
        mname = ws.cell(row=r, column=5).value
        lname = ws.cell(row=r, column=6).value
        email_val = ws.cell(row=r, column=8).value
        prog_val = ws.cell(row=r, column=2).value

        # Skip completely blank rows
        if not any([fname, mname, lname, email_val, prog_val]):
            continue

        parts = [str(p).strip() for p in (fname, mname, lname) if p and str(p).strip()]
        full_name = " ".join(parts)
        email = str(email_val).strip().lower() if email_val else ""
        program = str(prog_val).strip() if prog_val else ""

        if not email or "@" not in email:
            errors.append(schemas.BulkImportXlsxRowError(row=r, name=full_name, email=email, message="Invalid or missing email"))
            continue
        if not full_name:
            errors.append(schemas.BulkImportXlsxRowError(row=r, name=full_name, email=email, message="Missing candidate name"))
            continue
        if not program:
            errors.append(schemas.BulkImportXlsxRowError(row=r, name=full_name, email=email, message="Missing program"))
            continue

        domain = map_program(program)
        if domain == "Law":
            errors.append(schemas.BulkImportXlsxRowError(row=r, name=full_name, email=email, message=f"Unsupported program/domain: {program}"))
            continue
        if domain == "Unknown":
            errors.append(schemas.BulkImportXlsxRowError(row=r, name=full_name, email=email, message=f"Unknown program mapping: {program}"))
            continue

        existing = db.query(models.User).filter(models.User.email == email).first()
        if existing:
            already_existed += 1
            continue

        user = models.User(
            email=email,
            full_name=full_name,
            hashed_password=default_hash,
            role=models.UserRole.candidate,
            must_reset_password=True,
            is_active=True,
            approval_status="registered",
            registered_domain=domain
        )
        db.add(user)
        created += 1

    db.commit()
    return schemas.BulkImportXlsxResult(created=created, already_existed=already_existed, errors=errors)


@router.post("/admin/users/{user_id}/approve", response_model=schemas.ManagedUserOut)
def approve_candidate(
    user_id: int,
    db: Session = Depends(get_db),
    admin: models.User = Depends(security.get_current_admin)
):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user or user.role != models.UserRole.candidate:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Candidate not found")
    
    user.approval_status = "approved"
    user.token_version = (user.token_version or 0) + 1
    db.add(user)
    db.commit()
    db.refresh(user)
    # Inline serialization matching `_to_out` logic in users.py
    from backend.routers.users import _latest_session
    session = _latest_session(db, user)
    return schemas.ManagedUserOut(
        id=user.id, full_name=user.full_name, email=user.email, role=user.role.value,
        is_active=user.is_active, must_reset_password=user.must_reset_password,
        approval_status=user.approval_status, registered_domain=user.registered_domain,
        latest_session_status=session.status.value if session else None,
        latest_session_score=session.score if session else None
    )


@router.post("/admin/users/{user_id}/reject", response_model=schemas.ManagedUserOut)
def reject_candidate(
    user_id: int,
    db: Session = Depends(get_db),
    admin: models.User = Depends(security.get_current_admin)
):
    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user or user.role != models.UserRole.candidate:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Candidate not found")
    
    user.approval_status = "rejected"
    user.token_version = (user.token_version or 0) + 1
    db.add(user)
    db.commit()
    db.refresh(user)
    from backend.routers.users import _latest_session
    session = _latest_session(db, user)
    return schemas.ManagedUserOut(
        id=user.id, full_name=user.full_name, email=user.email, role=user.role.value,
        is_active=user.is_active, must_reset_password=user.must_reset_password,
        approval_status=user.approval_status, registered_domain=user.registered_domain,
        latest_session_status=session.status.value if session else None,
        latest_session_score=session.score if session else None
    )


@router.websocket("/ws/admin")
async def admin_ws(websocket: WebSocket, token: str = Query(...)):
    db = SessionLocal()
    try:
        security.decode_token_for_ws(token, db)
    except Exception:
        await websocket.close(code=4401)
        return
    finally:
        db.close()

    await manager.connect(websocket)
    try:
        while True:
            # JPEG snapshot feed and update messages are broadcast from candidates to admin
            await websocket.receive_text()
    except WebSocketDisconnect:
        await manager.disconnect(websocket)

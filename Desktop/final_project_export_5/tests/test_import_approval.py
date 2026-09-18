import pytest
import io
import openpyxl
from fastapi.testclient import TestClient
from backend.main import app
from backend.database import Base, engine, SessionLocal
from backend.models import User, UserRole
from backend import security

client = TestClient(app)

@pytest.fixture(scope="module")
def setup_database():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    # Create an admin user for testing
    admin_email = "admin_import_test@example.com"
    existing_admin = db.query(User).filter(User.email == admin_email).first()
    if not existing_admin:
        admin_user = User(
            email=admin_email,
            full_name="Admin Test",
            hashed_password=security.hash_password("adminpass"),
            role=UserRole.admin,
            is_active=True
        )
        db.add(admin_user)
        db.commit()
    db.close()
    yield
    # We won't drop all here as it might affect other tests, just rely on isolated data

@pytest.fixture
def admin_token(setup_database):
    response = client.post("/auth/login", json={"email": "admin_import_test@example.com", "password": "adminpass"})
    assert response.status_code == 200
    return response.json()["access_token"]


def test_xlsx_import(admin_token):
    # Create an in-memory XLSX file
    wb = openpyxl.Workbook()
    # Rename default sheet
    ws = wb.active
    ws.title = "Application Details"
    
    # Header row at 4
    headers = {2: "Program", 4: "First_Name", 5: "Middle_Name", 6: "Last_Name", 8: "Email_ID"}
    for col, name in headers.items():
        ws.cell(row=4, column=col).value = name
        
    # Data rows at 5+
    # Row 5: Valid User
    ws.cell(row=5, column=2).value = "M.Tech Computer Science" # CS
    ws.cell(row=5, column=4).value = "John"
    ws.cell(row=5, column=5).value = "H"
    ws.cell(row=5, column=6).value = "Doe"
    ws.cell(row=5, column=8).value = "john.import@example.com"

    # Row 6: Invalid Program (Law)
    ws.cell(row=6, column=2).value = "B.A. L.L.B Law" # Law -> should be rejected
    ws.cell(row=6, column=4).value = "Jane"
    ws.cell(row=6, column=6).value = "Smith"
    ws.cell(row=6, column=8).value = "jane.law@example.com"
    
    # Row 7: Blank First Name (but Last Name present) -> valid Full Name
    ws.cell(row=7, column=2).value = "M.Sc Mathematics" # Mathematics
    ws.cell(row=7, column=6).value = "Stark"
    ws.cell(row=7, column=8).value = "stark@example.com"
    
    # Row 8: Missing Email -> should be rejected
    ws.cell(row=8, column=2).value = "Tamil"
    ws.cell(row=8, column=4).value = "No"
    ws.cell(row=8, column=6).value = "Email"

    excel_bytes = io.BytesIO()
    wb.save(excel_bytes)
    excel_bytes.seek(0)
    
    headers_req = {"Authorization": f"Bearer {admin_token}"}
    files = {"file": ("test_import.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    
    response = client.post("/admin/users/upload-xlsx", headers=headers_req, files=files)
    assert response.status_code == 200
    data = response.json()
    
    assert data["created"] == 2, data # John and Stark
    assert data["already_existed"] == 0
    assert len(data["errors"]) == 2 # Law and Missing Email

def test_approval_workflow(admin_token):
    # Verify user was created with registered status
    db = SessionLocal()
    user = db.query(User).filter(User.email == "john.import@example.com").first()
    assert user is not None
    assert user.approval_status == "registered"
    assert user.registered_domain == "Computer Science & Engineering"
    user_id = user.id
    db.close()
    
    headers_req = {"Authorization": f"Bearer {admin_token}"}
    
    # 1. Reject
    res_reject = client.post(f"/admin/users/{user_id}/reject", headers=headers_req)
    assert res_reject.status_code == 200
    assert res_reject.json()["approval_status"] == "rejected"
    
    # 2. Approve
    res_approve = client.post(f"/admin/users/{user_id}/approve", headers=headers_req)
    assert res_approve.status_code == 200
    assert res_approve.json()["approval_status"] == "approved"

    # Verify candidate login and exam gating
    login_res = client.post("/auth/login", json={"email": "john.import@example.com", "password": "DefaultPass@2026"})
    assert login_res.status_code == 200
    candidate_token = login_res.json()["access_token"]
    assert login_res.json()["approval_status"] == "approved"
    
    # Must reset password first
    reset_res = client.post("/auth/reset-password", json={"new_password": "newpass123"}, headers={"Authorization": f"Bearer {candidate_token}"})
    assert reset_res.status_code == 200
    candidate_token = reset_res.json()["access_token"]
    
    # Start exam should now succeed in getting past the approval gate.
    # It will fail with 409 because the test DB has no questions, which is expected.
    exam_headers = {"Authorization": f"Bearer {candidate_token}"}
    start_res = client.post("/exam/start", headers=exam_headers)
    assert start_res.status_code == 409
    assert "Question bank isn't ready" in start_res.json()["detail"]
